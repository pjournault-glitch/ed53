import io
import unicodedata
from dataclasses import fields

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from model.model import Addressee, Animator, Participant, Sender

# libellés des en-têtes ; un champ absent de ce dictionnaire garde son nom technique
LABELS = {
    'title': 'Titre', 'firstname': 'Prénom', 'name': 'Nom', 'street': 'Adresse',
    'zip': 'CP', 'city': 'Ville', 'phone': 'Téléphone', 'mail': 'E-mail',
    'organisation': 'Organisation', 'sender_label': 'Envoyé par',
}
EXCLUDED = {'id', 'sender_id'}    # identifiants techniques, sans intérêt dans le fichier
MAX_WIDTH = 50                    # largeur maximale d'une colonne (en caractères)


def _sort_key(text) -> str:
    """Tri insensible à la casse et aux accents (« Étienne » avant « Zoé »)."""
    decomposed = unicodedata.normalize('NFD', text or '')
    return ''.join(c for c in decomposed if unicodedata.category(c) != 'Mn').casefold()


def _person_key(item) -> tuple:
    name = getattr(item, 'name', '') or getattr(item, 'organisation', '') or ''
    return _sort_key(name), _sort_key(getattr(item, 'firstname', ''))


class ExportExcelController:
    def __init__(self, participants, animators, addressees, senders):
        self.sheets = [
            ('Participants', Participant, participants),
            ('Animateurs', Animator, animators),
            ('Destinataires', Addressee, addressees),
            ('Envoyants', Sender, senders),
        ]

    def generate(self) -> bytes:
        wb = Workbook()
        wb.remove(wb.active)                      # supprime la feuille vide créée par défaut
        for title, cls, items in self.sheets:
            self._add_sheet(wb, title, cls, items)
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _columns(cls) -> list[str]:
        names = [f.name for f in fields(cls) if f.name not in EXCLUDED]
        if cls is Participant and 'sender_label' not in names:
            names.append('sender_label')          # l'envoyant d'un participant se lit par son libellé
        return names

    def _add_sheet(self, wb: Workbook, title: str, cls, items) -> None:
        ws = wb.create_sheet(title)
        columns = self._columns(cls)

        ws.append([LABELS.get(c, c) for c in columns])
        for item in sorted(items or [], key=_person_key):
            ws.append([self._value(item, c) for c in columns])

        # en-têtes : gras, fond gris, ligne figée, filtres automatiques
        header_fill = PatternFill('solid', fgColor='D9D9D9')
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(vertical='center')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

        # données : alignées en haut, retours à la ligne conservés (adresses sur 2 lignes)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)

        # largeur de chaque colonne d'après son contenu le plus long
        for idx in range(1, len(columns) + 1):
            letter = get_column_letter(idx)
            longest = max(
                (len(line) for cell in ws[letter] for line in str(cell.value or '').splitlines()),
                default=8,
            )
            ws.column_dimensions[letter].width = min(longest + 2, MAX_WIDTH)

    @staticmethod
    def _value(item, column: str):
        value = getattr(item, column, '')
        return '' if value is None else value